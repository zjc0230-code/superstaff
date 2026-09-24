from __future__ import annotations

import base64
from datetime import timedelta
from io import BytesIO

import pytest
from docx import Document
from sqlalchemy.pool import StaticPool
from sqlmodel import Session, SQLModel, create_engine, select

from app.agents.branching import ensure_open_gallery_binding
from app.api.knowledge import (
    confirm_discovery as confirm_discovery_api,
    list_documents,
    search_knowledge,
    update_chunk,
    update_document,
)
from app.api.knowledge_bases import knowledge_base_read
from app.db.models import (
    AgentProfile,
    KnowledgeBase,
    KnowledgeBaseVersion,
    KnowledgeBucket,
    KnowledgeChunk,
    KnowledgeConcept,
    KnowledgeDiscoverySuggestion,
    KnowledgeDocument,
    KnowledgeIngestJob,
    ModelConfig,
    Skill,
    Tenant,
    Tool,
    User,
    utc_now,
)
from app.knowledge.schema import KnowledgeChunkUpdateRequest, KnowledgeDocumentUpdateRequest, KnowledgeSearchRequest, KnowledgeSearchResponse
from app.knowledge.okf import search_concepts
from app.knowledge.parser import extract_text
from app.knowledge.service import (
    IngestPayload,
    KnowledgeDiscoveryConflictError,
    KnowledgeDiscoveryValidationError,
    KnowledgeService,
    _build_evidence_pack,
    _build_section_nodes,
    _chunk_text_related_groups,
    _document_card_for_route,
    _expand_related_chunks,
    _route_candidates,
    _score_documents,
    _score_text,
    _select_diverse_chunk_hits,
    validate_discovered_skill,
)
from app.llm import LLMClient
from app.observability.spans import bind_span_sink
from app.skills.skill_schema import SkillCard


def test_skill_card_rejects_legacy_steps_and_accepts_graph() -> None:
    with pytest.raises(Exception):
        SkillCard(
            skill_id="skill_test",
            name="测试技能",
            steps=[
                {
                    "step_id": "collect",
                    "name": "收集信息",
                    "instruction": "收集用户信息",
                    "expected_user_info": ["name"],
                    "allowed_actions": ["ask_user", "continue_flow"],
                }
            ],
        )

    card = SkillCard(
        skill_id="skill_test",
        name="测试技能",
        nodes=[
            {
                "node_id": "collect",
                "type": "collect_info",
                "name": "收集信息",
                "instruction": "收集用户信息",
                "expected_user_info": ["name"],
                "allowed_actions": ["ask_user", "continue_flow"],
            },
            {
                "node_id": "reply",
                "type": "response",
                "name": "回复",
                "instruction": "回复用户",
                "allowed_actions": ["answer_user"],
            },
        ],
        edges=[{"source_node_id": "collect", "next_node_id": "reply"}],
        start_node_id="collect",
        terminal_node_ids=["reply"],
    )

    assert card.start_node_id == "collect"
    assert card.terminal_node_ids == ["reply"]
    assert [node.node_id for node in card.nodes] == ["collect", "reply"]
    assert card.edges[0].source_node_id == "collect"
    assert card.edges[0].next_node_id == "reply"


def test_knowledge_ingest_creates_document_buckets_and_chunks_without_auto_discovery() -> None:
    with _test_session() as db:
        db.add(Tenant(id="tenant_demo", name="Demo"))
        db.add(KnowledgeBase(id="kb_demo", tenant_id="tenant_demo", name="默认知识库"))
        db.commit()
        service = KnowledgeService(db)
        job = service.create_ingest_job(
            IngestPayload(
                tenant_id="tenant_demo",
                knowledge_base_id="kb_demo",
                filename="policy.md",
                content_base64=_b64("# 售后政策\n用户可查询订单。\n\n# 配送\n根据地址评估配送。"),
            )
        )

        service._run_ingest_job(job.id)  # noqa: SLF001 - exercise persistent job logic synchronously.

        job = db.get(type(job), job.id)
        assert job is not None
        assert job.status == "succeeded"
        assert job.document_id
        document = db.get(KnowledgeDocument, job.document_id)
        assert document is not None
        assert document.metadata_json["document_card"]["title"]
        assert document.metadata_json["section_tree"]
        assert document.metadata_json["chunk_stats"]["total_chunks"] > 0
        assert document.metadata_json["bucket_quality"]
        buckets = db.exec(select(KnowledgeBucket).where(KnowledgeBucket.document_id == job.document_id)).all()
        assert buckets
        assert all(bucket.metadata_json.get("section_ids") for bucket in buckets)
        chunks = db.exec(select(KnowledgeChunk).where(KnowledgeChunk.document_id == job.document_id)).all()
        assert chunks
        assert all(chunk.metadata_json.get("section_path") for chunk in chunks)
        assert all(chunk.metadata_json.get("ingest_schema_version") == 2 for chunk in chunks)
        for chunk in chunks:
            metadata = chunk.metadata_json or {}
            if metadata.get("related_group_id"):
                assert chunk.id in metadata.get("related_chunk_ids", [])
                assert metadata.get("related_chunk_count", 0) > 1
            else:
                assert metadata.get("related_chunk_ids") == []
                assert metadata.get("related_chunk_count") == 0
        response = service.search(
            KnowledgeSearchRequest(
                tenant_id="tenant_demo",
                knowledge_base_ids=["kb_demo"],
                query="配送怎么处理",
                mode="debug",
                need_evidence_pack=True,
            )
        )
        phases = [item["phase"] for item in response.route_trace]
        assert "document_route" in phases
        assert "bucket_route" in phases
        assert "section_expand" in phases
        assert "evidence_pack" in phases
        assert response.selected_documents
        assert response.expanded_sections
        assert response.evidence_pack
        assert response.evidence_pack[0]["source_path"]
        assert response.evidence_pack[0]["excerpt"]
        assert response.chunks
        assert db.exec(select(KnowledgeDiscoverySuggestion)).all() == []


def test_related_chunk_expansion_returns_all_siblings_in_source_order() -> None:
    chunks = [
        KnowledgeChunk(
            id=f"chunk_{index}",
            tenant_id="tenant_demo",
            knowledge_base_id="kb_demo",
            document_id="doc_demo",
            bucket_id="bucket_demo",
            chunk_index=index,
            content=f"片段 {index}",
            metadata_json={"related_group_id": "doc_demo:bucket_demo:section_1"},
        )
        for index in range(3)
    ]

    expanded = _expand_related_chunks([chunks[1]], [chunks[2], chunks[1], chunks[0]])

    assert [chunk.id for chunk in expanded] == ["chunk_0", "chunk_1", "chunk_2"]
    evidence = _build_evidence_pack("片段 1", expanded)
    assert [item["chunk_id"] for item in evidence] == ["chunk_0", "chunk_1", "chunk_2"]


def test_related_chunk_expansion_is_bounded_around_hit() -> None:
    chunks = [
        KnowledgeChunk(
            id=f"chunk_{index}",
            tenant_id="tenant_demo",
            knowledge_base_id="kb_demo",
            document_id="doc_demo",
            bucket_id="bucket_demo",
            chunk_index=index,
            content=str(index) * 700,
            metadata_json={"related_group_id": "large_group"},
        )
        for index in range(12)
    ]

    expanded = _expand_related_chunks([chunks[6]], chunks)

    assert chunks[6] in expanded
    assert len(expanded) <= 6
    assert sum(len(chunk.content) for chunk in expanded) <= 4800
    assert [chunk.chunk_index for chunk in expanded] == sorted(
        chunk.chunk_index for chunk in expanded
    )


def test_related_groups_only_link_parts_from_one_oversized_paragraph() -> None:
    groups = _chunk_text_related_groups(
        "# 流程标题\n\n第一条独立说明。\n\n第二条独立说明。\n\n" + ("连续正文。" * 300),
        300,
    )

    assert all("流程标题" not in part for group in groups for part in group)
    assert len(groups[0]) == 1
    assert len(groups[-1]) > 1
    assert sum("第一条独立说明" in part for group in groups for part in group) == 1


def test_diverse_chunk_hits_cover_selected_buckets_before_filling() -> None:
    buckets = [
        KnowledgeBucket(
            id=f"bucket_{index}",
            tenant_id="tenant_demo",
            knowledge_base_id="kb_demo",
            document_id="doc_demo",
            bucket_key=f"bucket_{index}",
            title=f"主题 {index}",
            summary="",
        )
        for index in range(3)
    ]
    ranked = [
        KnowledgeChunk(
            id=chunk_id,
            tenant_id="tenant_demo",
            knowledge_base_id="kb_demo",
            document_id="doc_demo",
            bucket_id=bucket_id,
            chunk_index=index,
            content=chunk_id,
        )
        for index, (chunk_id, bucket_id) in enumerate(
            [
                ("chunk_a1", "bucket_0"),
                ("chunk_a2", "bucket_0"),
                ("chunk_b1", "bucket_1"),
                ("chunk_c1", "bucket_2"),
            ]
        )
    ]

    selected = _select_diverse_chunk_hits(ranked, buckets, 4)

    assert {chunk.bucket_id for chunk in selected} == {"bucket_0", "bucket_1", "bucket_2"}


def test_docx_headings_are_preserved_and_numbered_body_items_are_not_sections() -> None:
    document = Document()
    document.add_heading("HR 模块", level=1)
    document.add_heading("入离职办理", level=2)
    document.add_paragraph("1. 准备材料甲")
    document.add_paragraph("2. 提交登记表乙")
    payload = BytesIO()
    document.save(payload)

    text, file_type = extract_text("hr.docx", payload.getvalue())
    sections = _build_section_nodes(text)

    assert file_type == "docx"
    assert text.startswith("# HR 模块\n## 入离职办理")
    assert [section["title"] for section in sections] == ["HR 模块", "入离职办理"]
    assert "1. 准备材料甲" in sections[1]["content"]
    assert "2. 提交登记表乙" in sections[1]["content"]


def test_pdf_pages_receive_stable_section_labels(monkeypatch) -> None:
    class FakePage:
        def __init__(self, text: str):
            self.text = text

        def extract_text(self) -> str:
            return self.text

    class FakeReader:
        pages = [FakePage("First page policy"), FakePage(""), FakePage("Third page process")]

    monkeypatch.setattr("pypdf.PdfReader", lambda _stream: FakeReader())

    text, file_type = extract_text("policy.pdf", b"fake-pdf")
    sections = _build_section_nodes(text)

    assert file_type == "pdf"
    assert text.startswith("# PDF 文档\n\n## 第 1 页")
    assert [section["title"] for section in sections] == ["PDF 文档", "第 1 页", "第 3 页"]
    assert sections[-1]["path"] == "PDF 文档 / 第 3 页"


def test_long_section_continuations_share_one_related_section() -> None:
    paragraphs = [f"{index}. 流程内容" + ("说明" * 150) for index in range(1, 9)]

    sections = _build_section_nodes("# 办理流程\n\n" + "\n\n".join(paragraphs))

    assert len(sections) > 1
    assert {section["title"] for section in sections} == {"办理流程"}
    assert len({section["related_section_id"] for section in sections}) == 1


def test_knowledge_ingest_cancel_queued_job_clears_embedded_content() -> None:
    with _test_session() as db:
        db.add(Tenant(id="tenant_demo", name="Demo"))
        db.add(KnowledgeBase(id="kb_demo", tenant_id="tenant_demo", name="默认知识库"))
        db.commit()
        service = KnowledgeService(db)
        job = service.create_ingest_job(
            IngestPayload(
                tenant_id="tenant_demo",
                knowledge_base_id="kb_demo",
                filename="policy.md",
                content_base64=_b64("# 售后政策\n用户可查询订单。"),
            )
        )

        cancelled = service.cancel_ingest_job(job.id, "tenant_demo")

        assert cancelled is not None
        assert cancelled.status == "cancelled"
        assert cancelled.stage == "cancelled"
        assert cancelled.finished_at is not None
        assert cancelled.metadata_json["stage_label"] == "已取消"
        assert "content_base64" not in cancelled.metadata_json


def test_knowledge_ingest_cancel_running_job_cleans_partial_artifacts() -> None:
    with _test_session() as db:
        db.add(Tenant(id="tenant_demo", name="Demo"))
        db.add(KnowledgeBase(id="kb_demo", tenant_id="tenant_demo", name="默认知识库"))
        document = KnowledgeDocument(
            id="kdoc_partial",
            tenant_id="tenant_demo",
            knowledge_base_id="kb_demo",
            filename="partial.md",
            file_type="md",
            title="半成品",
            status="processing",
        )
        bucket = KnowledgeBucket(
            id="kbucket_partial",
            tenant_id="tenant_demo",
            knowledge_base_id="kb_demo",
            document_id=document.id,
            bucket_key="partial",
            title="半成品目录",
            summary="半成品摘要",
        )
        chunk = KnowledgeChunk(
            id="kchunk_partial",
            tenant_id="tenant_demo",
            knowledge_base_id="kb_demo",
            document_id=document.id,
            bucket_id=bucket.id,
            chunk_index=0,
            content="半成品引用",
        )
        concept = KnowledgeConcept(
            id="kconcept_partial",
            tenant_id="tenant_demo",
            knowledge_base_id="kb_demo",
            document_id=document.id,
            concept_id="partial",
            concept_type="Source Document",
            title="半成品概念",
            content_md="半成品概念",
        )
        suggestion = KnowledgeDiscoverySuggestion(
            id="kdisc_partial",
            tenant_id="tenant_demo",
            knowledge_base_id="kb_demo",
            document_id=document.id,
            bucket_id=bucket.id,
            suggestion_type="warning",
            title="半成品建议",
        )
        job = KnowledgeIngestJob(
            id="kjob_partial",
            tenant_id="tenant_demo",
            knowledge_base_id="kb_demo",
            document_id=document.id,
            filename="partial.md",
            status="running",
            stage="chunking",
            progress=0.62,
            metadata_json={"content_base64": _b64("partial"), "stage_label": "生成引用来源"},
        )
        db.add(document)
        db.add(bucket)
        db.add(chunk)
        db.add(concept)
        db.add(suggestion)
        db.add(job)
        db.commit()
        document_id = document.id
        bucket_id = bucket.id
        chunk_id = chunk.id
        concept_id = concept.id
        suggestion_id = suggestion.id
        service = KnowledgeService(db)

        cancelling = service.cancel_ingest_job(job.id, "tenant_demo")
        assert cancelling is not None
        assert cancelling.status == "cancel_requested"

        service._run_ingest_job(job.id)  # noqa: SLF001 - exercise persisted cancellation path.

        cancelled = db.get(KnowledgeIngestJob, job.id)
        assert cancelled is not None
        assert cancelled.status == "cancelled"
        assert cancelled.stage == "cancelled"
        assert cancelled.document_id is None
        assert cancelled.metadata_json["cancelled_document_id"] == document_id
        assert "content_base64" not in cancelled.metadata_json
        assert db.get(KnowledgeDocument, document_id) is None
        assert db.get(KnowledgeBucket, bucket_id) is None
        assert db.get(KnowledgeChunk, chunk_id) is None
        assert db.get(KnowledgeConcept, concept_id) is None
        assert db.get(KnowledgeDiscoverySuggestion, suggestion_id) is None


def test_knowledge_ingest_stale_cancel_request_finalizes_without_worker() -> None:
    with _test_session() as db:
        db.add(Tenant(id="tenant_demo", name="Demo"))
        db.add(KnowledgeBase(id="kb_demo", tenant_id="tenant_demo", name="默认知识库"))
        document = KnowledgeDocument(
            id="kdoc_stale_cancel",
            tenant_id="tenant_demo",
            knowledge_base_id="kb_demo",
            filename="stale.md",
            file_type="md",
            title="取消中的半成品",
            status="processing",
        )
        job = KnowledgeIngestJob(
            id="kjob_stale_cancel",
            tenant_id="tenant_demo",
            knowledge_base_id="kb_demo",
            document_id=document.id,
            filename="stale.md",
            status="cancel_requested",
            stage="chunking",
            progress=70.0,
            metadata_json={"content_base64": _b64("partial"), "stage_label": "取消中"},
            updated_at=utc_now() - timedelta(seconds=60),
        )
        db.add(document)
        db.add(job)
        db.commit()
        service = KnowledgeService(db)

        finalized = service.finalize_stale_cancel_requested_job(job)

        assert finalized is not None
        assert finalized.status == "cancelled"
        assert finalized.stage == "cancelled"
        assert finalized.document_id is None
        assert finalized.metadata_json["stage_label"] == "已取消"
        assert "content_base64" not in finalized.metadata_json
        assert db.get(KnowledgeDocument, document.id) is None


def test_knowledge_search_without_model_uses_relevance_rank_order() -> None:
    with _test_session() as db:
        db.add(Tenant(id="tenant_demo", name="Demo"))
        db.add(KnowledgeBase(id="kb_demo", tenant_id="tenant_demo", name="默认知识库"))
        document = KnowledgeDocument(
            id="kdoc_frontend",
            tenant_id="tenant_demo",
            knowledge_base_id="kb_demo",
            filename="frontend.md",
            file_type="md",
            title="前端规范资料",
            status="ready",
            bucket_count=2,
            chunk_count=2,
            metadata_json={
                "document_card": {
                    "title": "前端规范资料",
                    "summary": "前端编码规范、Vue 3、组件规范和命名规范。",
                }
            },
        )
        irrelevant = KnowledgeBucket(
            id="kbucket_citation",
            tenant_id="tenant_demo",
            knowledge_base_id="kb_demo",
            document_id=document.id,
            bucket_key="citation",
            title="知识引用测试说明",
            summary="回答引用展示规则。",
        )
        frontend = KnowledgeBucket(
            id="kbucket_frontend",
            tenant_id="tenant_demo",
            knowledge_base_id="kb_demo",
            document_id=document.id,
            bucket_key="frontend",
            title="前端编码规范",
            summary="Vue 3、Vite、TypeScript、组件编写和命名规范。",
        )
        db.add(document)
        db.add(irrelevant)
        db.add(frontend)
        db.add(
            KnowledgeChunk(
                tenant_id="tenant_demo",
                knowledge_base_id="kb_demo",
                document_id=document.id,
                bucket_id=irrelevant.id,
                chunk_index=0,
                content="知识引用展示规则。",
                summary="知识引用展示规则。",
                source_ref="citation.md",
            )
        )
        db.add(
            KnowledgeChunk(
                tenant_id="tenant_demo",
                knowledge_base_id="kb_demo",
                document_id=document.id,
                bucket_id=frontend.id,
                chunk_index=0,
                content="前端规范包括 Vue 3、Vite、TypeScript 和组件编写规范。",
                summary="前端规范包括 Vue 3、Vite、TypeScript 和组件编写规范。",
                source_ref="frontend.md",
            )
        )
        db.commit()

        response = KnowledgeService(db).search(
            KnowledgeSearchRequest(
                tenant_id="tenant_demo",
                knowledge_base_ids=["kb_demo"],
                query="前端规范有哪些？",
                mode="chat",
                max_buckets=2,
                need_evidence_pack=True,
            )
        )

        assert [bucket.id for bucket in response.selected_buckets] == ["kbucket_frontend"]


def test_model_driven_document_route_does_not_fall_back_to_lexical_matching(monkeypatch) -> None:
    with _test_session() as db:
        db.add(Tenant(id="tenant_demo", name="Demo"))
        db.add(KnowledgeBase(id="kb_demo", tenant_id="tenant_demo", name="默认知识库"))
        db.add(
            KnowledgeDocument(
                id="kdoc_frontend",
                tenant_id="tenant_demo",
                knowledge_base_id="kb_demo",
                filename="frontend.md",
                file_type="md",
                title="前端规范资料",
                status="ready",
                metadata_json={"document_card": {"title": "前端规范资料", "summary": "前端编码规范。"}},
            )
        )
        db.commit()
        monkeypatch.setattr(KnowledgeService, "_select_documents_with_llm", lambda *args, **kwargs: [])

        response = KnowledgeService(db).search(
            KnowledgeSearchRequest(
                tenant_id="tenant_demo",
                knowledge_base_ids=["kb_demo"],
                query="前端规范有哪些？",
                mode="chat",
            ),
            ModelConfig(id="model_route", tenant_id="tenant_demo", name="Route", model="route"),
        )

        assert response.selected_documents == []
        assert any(item.get("phase") == "document_route_no_match" for item in response.route_trace)
        assert all("fallback" not in str(item.get("phase") or "") for item in response.route_trace)


def test_model_route_failure_falls_back_to_lexical_matching(monkeypatch) -> None:
    with _test_session() as db:
        db.add(Tenant(id="tenant_demo", name="Demo"))
        db.add(KnowledgeBase(id="kb_demo", tenant_id="tenant_demo", name="默认知识库"))
        document = KnowledgeDocument(
            id="kdoc_leave",
            tenant_id="tenant_demo",
            knowledge_base_id="kb_demo",
            filename="leave.md",
            file_type="md",
            title="离职办理",
            status="ready",
            metadata_json={"document_card": {"title": "离职办理", "summary": "离职申请流程"}},
        )
        bucket = KnowledgeBucket(
            id="kbucket_leave",
            tenant_id="tenant_demo",
            knowledge_base_id="kb_demo",
            document_id=document.id,
            bucket_key="leave",
            title="离职申请流程",
            summary="员工申请离职需要提交审批",
        )
        db.add(document)
        db.add(bucket)
        db.add(
            KnowledgeChunk(
                tenant_id="tenant_demo",
                knowledge_base_id="kb_demo",
                document_id=document.id,
                bucket_id=bucket.id,
                chunk_index=0,
                content="员工申请离职后进入标准审批流程。",
            )
        )
        db.commit()
        monkeypatch.setattr(
            KnowledgeService, "_select_documents_with_llm", lambda *args, **kwargs: None
        )
        monkeypatch.setattr(
            KnowledgeService, "_select_buckets_with_llm", lambda *args, **kwargs: None
        )

        response = KnowledgeService(db).search(
            KnowledgeSearchRequest(
                tenant_id="tenant_demo",
                knowledge_base_ids=["kb_demo"],
                query="怎么申请离职",
            ),
            ModelConfig(id="model_route", tenant_id="tenant_demo", name="Route", model="route"),
        )

        phases = {item.get("phase") for item in response.route_trace}
        assert response.chunks
        assert "document_route_lexical_fallback" in phases
        assert "bucket_route_lexical_fallback" in phases


def test_document_loading_does_not_hide_relevant_rows_after_first_40() -> None:
    with _test_session() as db:
        db.add(Tenant(id="tenant_demo", name="Demo"))
        db.add(KnowledgeBase(id="kb_demo", tenant_id="tenant_demo", name="默认知识库"))
        relevant = KnowledgeDocument(
            id="kdoc_relevant_old",
            tenant_id="tenant_demo",
            knowledge_base_id="kb_demo",
            filename="social-security.md",
            file_type="md",
            title="社保公积金办理",
            status="ready",
            metadata_json={"document_card": {"summary": "医保与社保办理规则"}},
        )
        db.add(relevant)
        for index in range(145):
            db.add(
                KnowledgeDocument(
                    id=f"kdoc_irrelevant_{index}",
                    tenant_id="tenant_demo",
                    knowledge_base_id="kb_demo",
                    filename=f"other-{index}.md",
                    file_type="md",
                    title=f"其他资料 {index}",
                    status="ready",
                )
            )
        db.commit()

        documents = KnowledgeService(db)._load_documents_for_search(
            KnowledgeSearchRequest(
                tenant_id="tenant_demo",
                knowledge_base_ids=["kb_demo"],
                query="社保怎么办理",
            )
        )

        assert len(documents) == 146
        assert relevant.id in {document.id for document in documents}
        route_documents = _route_candidates(
            _score_documents("社保怎么办理", documents), documents, 120
        )
        assert relevant.id in {document.id for document in route_documents}


def test_document_route_card_keeps_late_outline_entries() -> None:
    document = KnowledgeDocument(
        id="kdoc_hr",
        tenant_id="tenant_demo",
        knowledge_base_id="kb_demo",
        filename="hr.md",
        file_type="md",
        title="HR FAQ",
        status="ready",
        metadata_json={
            "document_card": {
                "outline": [
                    {"title": f"章节 {index}", "path": f"HR / 章节 {index}"}
                    for index in range(1, 17)
                ]
            }
        },
    )

    card = _document_card_for_route(document)

    assert card["outline"][-1] == "HR / 章节 16"


def test_query_scoring_reduces_generic_question_noise() -> None:
    query = "怎么申请离职"

    assert _score_text(query, "员工申请离职后进入标准审批流程") > _score_text(
        query, "申请在职收入证明"
    )


def test_knowledge_search_records_persistent_substep_spans() -> None:
    events: list[tuple[str, dict]] = []
    with _test_session() as db:
        db.add(Tenant(id="tenant_demo", name="Demo"))
        db.add(KnowledgeBase(id="kb_demo", tenant_id="tenant_demo", name="默认知识库"))
        document = KnowledgeDocument(
            id="kdoc_frontend",
            tenant_id="tenant_demo",
            knowledge_base_id="kb_demo",
            filename="frontend.md",
            file_type="md",
            title="前端规范资料",
            status="ready",
            metadata_json={"document_card": {"title": "前端规范资料", "summary": "前端规范"}},
        )
        bucket = KnowledgeBucket(
            id="kbucket_frontend",
            tenant_id="tenant_demo",
            knowledge_base_id="kb_demo",
            document_id=document.id,
            bucket_key="frontend",
            title="前端规范",
            summary="Vue 3 与 TypeScript",
        )
        db.add(document)
        db.add(bucket)
        db.add(
            KnowledgeChunk(
                tenant_id="tenant_demo",
                knowledge_base_id="kb_demo",
                document_id=document.id,
                bucket_id=bucket.id,
                chunk_index=0,
                content="前端规范包括 Vue 3 与 TypeScript。",
                summary="前端规范",
                source_ref="frontend.md",
            )
        )
        db.commit()

        with bind_span_sink(
            lambda event_type, payload: events.append((event_type, payload))
        ):
            response = KnowledgeService(db).search(
                KnowledgeSearchRequest(
                    tenant_id="tenant_demo",
                    knowledge_base_ids=["kb_demo"],
                    query="前端规范",
                    mode="chat",
                    need_evidence_pack=True,
                )
            )

    assert response.chunks
    finished = {
        payload["operation"]: payload
        for event_type, payload in events
        if event_type == "knowledge_span_finished"
    }
    assert {
        "knowledge.search",
        "knowledge.load_concepts",
        "knowledge.route_concepts",
        "knowledge.load_documents",
        "knowledge.route_documents",
        "knowledge.load_buckets",
        "knowledge.route_buckets",
        "knowledge.expand_sections",
        "knowledge.load_chunks",
        "knowledge.rank_chunks",
        "knowledge.build_evidence_pack",
    }.issubset(finished)
    assert finished["knowledge.search"]["duration_ms"] >= 0
    assert finished["knowledge.load_documents"]["candidate_count"] == 1
    assert finished["knowledge.build_evidence_pack"]["evidence_count"] == 1


def test_legacy_knowledge_without_new_metadata_remains_searchable() -> None:
    with _test_session() as db:
        db.add(Tenant(id="tenant_demo", name="Demo"))
        db.add(KnowledgeBase(id="kb_demo", tenant_id="tenant_demo", name="默认知识库"))
        document = KnowledgeDocument(
            id="kdoc_legacy",
            tenant_id="tenant_demo",
            knowledge_base_id="kb_demo",
            filename="legacy-policy.txt",
            file_type="txt",
            title="年假制度",
            status="ready",
            metadata_json={},
        )
        bucket = KnowledgeBucket(
            id="kbucket_legacy",
            tenant_id="tenant_demo",
            knowledge_base_id="kb_demo",
            document_id=document.id,
            bucket_key="legacy",
            title="年假制度",
            summary="",
            metadata_json={},
        )
        chunk = KnowledgeChunk(
            id="kchunk_legacy",
            tenant_id="tenant_demo",
            knowledge_base_id="kb_demo",
            document_id=document.id,
            bucket_id=bucket.id,
            chunk_index=0,
            content="年假余额以系统中的有效记录为准。",
            metadata_json={},
        )
        db.add(document)
        db.add(bucket)
        db.add(chunk)
        db.commit()

        response = KnowledgeService(db).search(
            KnowledgeSearchRequest(
                tenant_id="tenant_demo",
                knowledge_base_ids=["kb_demo"],
                query="年假",
            )
        )

        assert [item.id for item in response.chunks] == ["kchunk_legacy"]
        assert response.evidence_pack[0]["related_group_id"] is None


def test_okf_search_does_not_require_manually_curated_business_terms() -> None:
    concept = KnowledgeConcept(
        tenant_id="tenant_demo",
        knowledge_base_id="kb_demo",
        concept_id="sources/internal-document",
        concept_type="Source Document",
        title="内部文档说明",
        description="介绍可用文档及其适用范围。",
        content_md="# 内部文档说明\n\n这份文档记录服务流程。",
    )

    assert search_concepts("文档", [concept]) == [concept]


def test_knowledge_search_api_uses_selected_model_config(monkeypatch) -> None:
    captured: dict[str, str | None] = {}

    def fake_search(self, request, model_config=None):  # noqa: ANN001
        captured["model_id"] = model_config.id if model_config else None
        return KnowledgeSearchResponse(route_trace=[{"phase": "ok"}])

    monkeypatch.setattr(KnowledgeService, "search", fake_search)

    with _test_session() as db:
        db.add(Tenant(id="tenant_demo", name="Demo"))
        db.add(AgentProfile(id="agent_overall", tenant_id="tenant_demo", name="开放广场", is_overall=True))
        db.add(KnowledgeBase(id="kb_search", tenant_id="tenant_demo", name="检索知识库"))
        db.add(
            ModelConfig(
                id="model_default",
                tenant_id="tenant_demo",
                name="Default model",
                api_key_encrypted="",
                model="default",
                is_default=True,
                enabled=True,
            )
        )
        db.add(
            ModelConfig(
                id="model_selected",
                tenant_id="tenant_demo",
                name="Selected model",
                api_key_encrypted="",
                model="selected",
                enabled=True,
            )
        )
        ensure_open_gallery_binding(db, "tenant_demo", "knowledge_base", "kb_search", "active")
        db.commit()

        search_knowledge(
            KnowledgeSearchRequest(
                tenant_id="tenant_demo",
                query="测试检索",
                model_config_id="model_selected",
            ),
            db,
            User(id="user_admin", tenant_id="tenant_demo", username="admin", role="admin"),
        )

        assert captured["model_id"] == "model_selected"


def test_knowledge_base_read_keeps_archived_rows_visible_despite_active_versions() -> None:
    row = KnowledgeBase(id="kb_demo", tenant_id="tenant_demo", name="默认知识库", status="archived")
    version = KnowledgeBaseVersion(
        tenant_id="tenant_demo",
        knowledge_base_id=row.id,
        version="1.0.0",
        name=row.name,
        status="active",
    )

    overall_read = knowledge_base_read(row, {}, version_row=version)
    branch_read = knowledge_base_read(
        row,
        {},
        version_row=version,
        branch_meta={"status": "inactive", "base_version": "1.0.0", "head_version": "1.0.0", "sync_state": "synced"},
    )

    assert overall_read.status == "archived"
    assert branch_read.status == "archived"


def test_list_documents_without_agent_scope_returns_only_open_gallery_documents() -> None:
    with _test_session() as db:
        db.add(Tenant(id="tenant_demo", name="Demo"))
        db.add(AgentProfile(id="agent_overall", tenant_id="tenant_demo", name="开放广场", is_overall=True))
        db.add(KnowledgeBase(id="kb_open", tenant_id="tenant_demo", name="开放知识库"))
        db.add(KnowledgeBase(id="kb_private", tenant_id="tenant_demo", name="私有知识库"))
        db.add(
            KnowledgeBaseVersion(
                id="kbv_open",
                tenant_id="tenant_demo",
                knowledge_base_id="kb_open",
                version="1.0.0",
                name="开放知识库",
            )
        )
        db.add(
            KnowledgeBaseVersion(
                id="kbv_private",
                tenant_id="tenant_demo",
                knowledge_base_id="kb_private",
                version="1.0.0",
                name="私有知识库",
            )
        )
        db.add(
            KnowledgeDocument(
                id="kdoc_open",
                tenant_id="tenant_demo",
                knowledge_base_id="kb_open",
                knowledge_base_version_id="kbv_open",
                filename="open.md",
                file_type="md",
                title="开放资料",
                status="ready",
            )
        )
        db.add(
            KnowledgeDocument(
                id="kdoc_private",
                tenant_id="tenant_demo",
                knowledge_base_id="kb_private",
                knowledge_base_version_id="kbv_private",
                filename="private.md",
                file_type="md",
                title="私有资料",
                status="ready",
            )
        )
        db.flush()
        ensure_open_gallery_binding(db, "tenant_demo", "knowledge_base", "kb_open", "active")
        db.commit()

        rows = list_documents("tenant_demo", None, None, True, db)

        assert {row.id for row in rows} == {"kdoc_open"}


def test_update_document_syncs_document_card_and_okf_source_concept() -> None:
    with _test_session() as db:
        db.add(Tenant(id="tenant_demo", name="Demo"))
        db.add(KnowledgeBase(id="kb_demo", tenant_id="tenant_demo", name="默认知识库"))
        db.add(
            KnowledgeBaseVersion(
                id="kbv_demo",
                tenant_id="tenant_demo",
                knowledge_base_id="kb_demo",
                version="1.0.0",
                name="默认知识库",
                status="active",
            )
        )
        document = KnowledgeDocument(
            id="kdoc_demo",
            tenant_id="tenant_demo",
            knowledge_base_id="kb_demo",
            knowledge_base_version_id="kbv_demo",
            filename="demo.md",
            file_type="md",
            title="旧标题",
            status="ready",
            bucket_count=1,
            chunk_count=1,
            metadata_json={
                "document_card": {"title": "旧卡片标题", "summary": "文档摘要"},
                "section_tree": [
                    {
                        "section_id": "intro",
                        "title": "介绍",
                        "path": "介绍",
                        "summary": "旧章节摘要",
                        "content": "旧章节内容",
                    }
                ],
            },
        )
        bucket = KnowledgeBucket(
            tenant_id="tenant_demo",
            knowledge_base_id="kb_demo",
            knowledge_base_version_id="kbv_demo",
            document_id=document.id,
            bucket_key="intro",
            title="介绍",
            summary="旧桶摘要",
            token_estimate=10,
            metadata_json={"content": "旧桶内容", "section_ids": ["intro"], "section_paths": ["介绍"]},
        )
        stale_source = KnowledgeConcept(
            tenant_id="tenant_demo",
            knowledge_base_id="kb_demo",
            knowledge_base_version_id="kbv_demo",
            document_id=document.id,
            concept_id="sources/old-title",
            concept_type="Source Document",
            title="旧卡片标题",
            description="旧来源",
            content_md="# Old",
        )
        db.add(document)
        db.add(bucket)
        db.add(stale_source)
        db.commit()

        updated = update_document(
            document.id,
            KnowledgeDocumentUpdateRequest(tenant_id="tenant_demo", title="新标题"),
            db,
        )

        assert updated.title == "新标题"
        assert updated.metadata["document_card"]["title"] == "新标题"
        source_concepts = db.exec(
            select(KnowledgeConcept).where(
                KnowledgeConcept.tenant_id == "tenant_demo",
                KnowledgeConcept.document_id == document.id,
                KnowledgeConcept.concept_type == "Source Document",
            )
        ).all()
        assert len(source_concepts) == 1
        assert source_concepts[0].title == "新标题"
        assert source_concepts[0].concept_id != "sources/old-title"


def test_update_document_content_rebuilds_all_derived_knowledge() -> None:
    with _test_session() as db:
        db.add(Tenant(id="tenant_demo", name="Demo"))
        db.add(KnowledgeBase(id="kb_demo", tenant_id="tenant_demo", name="默认知识库"))
        db.add(
            KnowledgeBaseVersion(
                id="kbv_demo",
                tenant_id="tenant_demo",
                knowledge_base_id="kb_demo",
                version="1.0.0",
                name="默认知识库",
                status="active",
            )
        )
        document = KnowledgeDocument(
            id="kdoc_demo",
            tenant_id="tenant_demo",
            knowledge_base_id="kb_demo",
            knowledge_base_version_id="kbv_demo",
            filename="demo.md",
            file_type="md",
            title="旧标题",
            status="ready",
            bucket_count=1,
            chunk_count=1,
            metadata_json={"raw_text": "# 旧内容\n\n即将被替换。"},
        )
        bucket = KnowledgeBucket(
            id="kbucket_old",
            tenant_id="tenant_demo",
            knowledge_base_id="kb_demo",
            knowledge_base_version_id="kbv_demo",
            document_id=document.id,
            bucket_key="old",
            title="旧索引",
            summary="旧摘要",
            token_estimate=10,
            metadata_json={"content": "旧内容"},
        )
        db.add(document)
        db.add(bucket)
        db.add(
            KnowledgeChunk(
                tenant_id="tenant_demo",
                knowledge_base_id="kb_demo",
                knowledge_base_version_id="kbv_demo",
                document_id=document.id,
                bucket_id=bucket.id,
                chunk_index=0,
                content="旧内容",
            )
        )
        db.add(
            KnowledgeConcept(
                tenant_id="tenant_demo",
                knowledge_base_id="kb_demo",
                knowledge_base_version_id="kbv_demo",
                document_id=document.id,
                concept_id="sources/old-title",
                concept_type="Source Document",
                title="旧标题",
                content_md="# 旧标题\n\n旧内容",
            )
        )
        db.commit()

        updated = update_document(
            document.id,
            KnowledgeDocumentUpdateRequest(
                tenant_id="tenant_demo",
                title="新版制度",
                content_md="# 新版制度\n\n## 适用范围\n\n仅适用于在线编辑测试。",
                expected_updated_at=document.updated_at.isoformat(),
            ),
            db,
        )

        assert updated.title == "新版制度"
        assert updated.metadata["raw_text"].startswith("# 新版制度")
        assert updated.metadata["section_stats"]["section_count"] >= 2
        assert updated.bucket_count >= 1
        assert updated.chunk_count >= 1
        chunks = db.exec(
            select(KnowledgeChunk).where(KnowledgeChunk.document_id == document.id)
        ).all()
        assert chunks
        assert all("旧内容" not in row.content for row in chunks)
        assert any("在线编辑测试" in row.content for row in chunks)
        concepts = db.exec(
            select(KnowledgeConcept).where(KnowledgeConcept.document_id == document.id)
        ).all()
        assert concepts
        assert all(row.concept_id != "sources/old-title" for row in concepts)
        assert any(row.title == "新版制度" for row in concepts)


def test_update_document_content_rejects_stale_editor_revision() -> None:
    with _test_session() as db:
        db.add(Tenant(id="tenant_demo", name="Demo"))
        db.add(KnowledgeBase(id="kb_demo", tenant_id="tenant_demo", name="默认知识库"))
        document = KnowledgeDocument(
            id="kdoc_demo",
            tenant_id="tenant_demo",
            knowledge_base_id="kb_demo",
            filename="demo.md",
            file_type="md",
            title="制度",
            status="ready",
        )
        db.add(document)
        db.commit()

        with pytest.raises(Exception) as exc_info:
            update_document(
                document.id,
                KnowledgeDocumentUpdateRequest(
                    tenant_id="tenant_demo",
                    content_md="# 冲突内容",
                    expected_updated_at="2020-01-01T00:00:00",
                ),
                db,
            )

        assert getattr(exc_info.value, "status_code", None) == 409


def test_update_chunk_refreshes_bucket_content_and_okf_topic() -> None:
    with _test_session() as db:
        db.add(Tenant(id="tenant_demo", name="Demo"))
        db.add(KnowledgeBase(id="kb_demo", tenant_id="tenant_demo", name="默认知识库"))
        db.add(
            KnowledgeBaseVersion(
                id="kbv_demo",
                tenant_id="tenant_demo",
                knowledge_base_id="kb_demo",
                version="1.0.0",
                name="默认知识库",
                status="active",
            )
        )
        document = KnowledgeDocument(
            id="kdoc_demo",
            tenant_id="tenant_demo",
            knowledge_base_id="kb_demo",
            knowledge_base_version_id="kbv_demo",
            filename="demo.md",
            file_type="md",
            title="测试文档",
            status="ready",
            bucket_count=1,
            chunk_count=1,
            metadata_json={"document_card": {"title": "测试文档", "summary": "文档摘要"}},
        )
        bucket = KnowledgeBucket(
            id="kbucket_demo",
            tenant_id="tenant_demo",
            knowledge_base_id="kb_demo",
            knowledge_base_version_id="kbv_demo",
            document_id=document.id,
            bucket_key="refund",
            title="退款规则",
            summary="旧退款规则摘要",
            token_estimate=10,
            metadata_json={"content": "旧退款规则内容"},
        )
        chunk = KnowledgeChunk(
            id="kchunk_demo",
            tenant_id="tenant_demo",
            knowledge_base_id="kb_demo",
            knowledge_base_version_id="kbv_demo",
            document_id=document.id,
            bucket_id=bucket.id,
            chunk_index=0,
            content="旧退款规则内容",
            summary="旧摘要",
        )
        db.add(document)
        db.add(bucket)
        db.add(chunk)
        db.commit()

        update_chunk(
            chunk.id,
            KnowledgeChunkUpdateRequest(tenant_id="tenant_demo", content="新退款规则内容", summary="新摘要"),
            db,
        )

        refreshed_bucket = db.get(KnowledgeBucket, bucket.id)
        assert refreshed_bucket is not None
        assert "新退款规则内容" in refreshed_bucket.metadata_json["content"]
        topic = db.exec(
            select(KnowledgeConcept).where(
                KnowledgeConcept.tenant_id == "tenant_demo",
                KnowledgeConcept.document_id == document.id,
                KnowledgeConcept.title == "退款规则",
            )
        ).one()
        assert "新退款规则内容" in topic.content_md


def test_confirm_discovery_is_required_before_tool_or_skill_enters_runtime() -> None:
    with _test_session() as db:
        db.add(Tenant(id="tenant_demo", name="Demo"))
        db.add(KnowledgeBase(id="kb_demo", tenant_id="tenant_demo", name="默认知识库"))
        suggestion = KnowledgeDiscoverySuggestion(
            tenant_id="tenant_demo",
            knowledge_base_id="kb_demo",
            document_id="doc_1",
            suggestion_type="tool",
            title="会员权益核对",
            payload_json={
                "name": "member.benefit_reconcile",
                "display_name": "会员权益核对",
                "method": "POST",
                "url": "/api/mock/member/benefit-reconcile",
            },
        )
        db.add(suggestion)
        db.commit()
        db.refresh(suggestion)

        assert db.exec(select(Tool)).all() == []
        result = KnowledgeService(db).confirm_discovery(suggestion)

        assert result["status"] == "created"
        assert db.exec(select(Tool).where(Tool.name == "member.benefit_reconcile")).first()
        assert db.exec(select(Skill)).all() == []


def test_confirm_discovery_rejects_tool_without_url() -> None:
    with _test_session() as db:
        db.add(Tenant(id="tenant_demo", name="Demo"))
        db.add(KnowledgeBase(id="kb_demo", tenant_id="tenant_demo", name="默认知识库"))
        suggestion = KnowledgeDiscoverySuggestion(
            tenant_id="tenant_demo",
            knowledge_base_id="kb_demo",
            document_id="doc_1",
            suggestion_type="tool",
            title="缺少地址的工具",
            payload_json={"name": "missing.url"},
        )
        db.add(suggestion)
        db.commit()

        with pytest.raises(KnowledgeDiscoveryValidationError, match="缺少 url"):
            KnowledgeService(db).confirm_discovery(suggestion)

        db.refresh(suggestion)
        assert suggestion.status == "pending"
        assert db.exec(select(Tool)).all() == []


def test_confirm_discovery_api_returns_422_for_invalid_skill() -> None:
    with _test_session() as db:
        db.add(Tenant(id="tenant_demo", name="Demo"))
        db.add(KnowledgeBase(id="kb_demo", tenant_id="tenant_demo", name="默认知识库"))
        user = User(
            id="user_admin",
            tenant_id="tenant_demo",
            username="admin",
            role="admin",
            password_hash="unused",
        )
        suggestion = KnowledgeDiscoverySuggestion(
            tenant_id="tenant_demo",
            knowledge_base_id="kb_demo",
            document_id="doc_1",
            suggestion_type="skill",
            title="错误格式的技能",
            payload_json={"draft_skill": {"skill_id": "invalid", "name": "错误格式"}},
        )
        db.add(user)
        db.add(suggestion)
        db.commit()

        with pytest.raises(Exception) as exc_info:
            confirm_discovery_api(suggestion.id, "tenant_demo", db, user)

        assert getattr(exc_info.value, "status_code", None) == 422
        assert "SuperStaff SkillCard" in str(getattr(exc_info.value, "detail", ""))


def test_confirm_discovery_api_returns_409_for_non_pending_status() -> None:
    with _test_session() as db:
        db.add(Tenant(id="tenant_demo", name="Demo"))
        db.add(KnowledgeBase(id="kb_demo", tenant_id="tenant_demo", name="默认知识库"))
        user = User(
            id="user_admin",
            tenant_id="tenant_demo",
            username="admin",
            role="admin",
            password_hash="unused",
        )
        suggestion = KnowledgeDiscoverySuggestion(
            tenant_id="tenant_demo",
            knowledge_base_id="kb_demo",
            document_id="doc_1",
            suggestion_type="warning",
            title="已处理建议",
            status="confirmed",
        )
        db.add(user)
        db.add(suggestion)
        db.commit()

        with pytest.raises(Exception) as exc_info:
            confirm_discovery_api(suggestion.id, "tenant_demo", db, user)

        assert getattr(exc_info.value, "status_code", None) == 409
        assert "只有待处理建议可以确认" in str(getattr(exc_info.value, "detail", ""))


def test_confirm_discovery_rejects_noncanonical_skill_graph() -> None:
    with _test_session() as db:
        db.add(Tenant(id="tenant_demo", name="Demo"))
        db.add(KnowledgeBase(id="kb_demo", tenant_id="tenant_demo", name="默认知识库"))
        suggestion = KnowledgeDiscoverySuggestion(
            tenant_id="tenant_demo",
            knowledge_base_id="kb_demo",
            document_id="doc_1",
            suggestion_type="skill",
            title="差旅报销审批",
            payload_json={
                "draft_skill": {
                    "skill_id": "expense.travel_approval",
                    "name": "差旅报销审批",
                    "nodes": [
                        {"id": "start", "type": "start", "label": "开始"},
                        {"id": "approve", "type": "action", "label": "主管审批", "description": "核对单据"},
                        {"id": "end", "type": "terminal", "label": "完成"},
                    ],
                    "edges": [
                        {"from": "start", "to": "approve"},
                        {"source": "approve", "target": "end"},
                    ],
                    "start_node_id": "start",
                    "terminal_node_ids": ["end"],
                }
            },
        )
        db.add(suggestion)
        db.commit()
        db.refresh(suggestion)

        with pytest.raises(KnowledgeDiscoveryValidationError, match="nodes.0.*id, label"):
            KnowledgeService(db).confirm_discovery(suggestion)

        db.refresh(suggestion)
        assert suggestion.status == "pending"
        assert db.exec(select(Skill)).all() == []


def test_confirm_discovery_does_not_overwrite_existing_skill() -> None:
    card = SkillCard(
        skill_id="expense.travel_approval",
        name="知识发现技能",
        nodes=[
            {
                "node_id": "reply",
                "type": "response",
                "name": "反馈结果",
                "instruction": "反馈结果。",
            }
        ],
        edges=[],
        start_node_id="reply",
        terminal_node_ids=["reply"],
    )
    with _test_session() as db:
        db.add(Tenant(id="tenant_demo", name="Demo"))
        db.add(KnowledgeBase(id="kb_demo", tenant_id="tenant_demo", name="默认知识库"))
        existing = Skill(
            tenant_id="tenant_demo",
            skill_id=card.skill_id,
            name="生产技能",
            version="2.0.0",
            content_json={**card.model_dump(mode="json"), "name": "生产技能", "version": "2.0.0"},
            status="published",
        )
        suggestion = KnowledgeDiscoverySuggestion(
            tenant_id="tenant_demo",
            knowledge_base_id="kb_demo",
            document_id="doc_1",
            suggestion_type="skill",
            title=card.name,
            payload_json={"draft_skill": card.model_dump(mode="json")},
        )
        db.add(existing)
        db.add(suggestion)
        db.commit()

        with pytest.raises(KnowledgeDiscoveryConflictError, match="不能通过知识发现覆盖"):
            KnowledgeService(db).confirm_discovery(suggestion)

        db.refresh(existing)
        db.refresh(suggestion)
        assert existing.name == "生产技能"
        assert existing.version == "2.0.0"
        assert existing.status == "published"
        assert suggestion.status == "pending"


@pytest.mark.parametrize("status", ["confirmed", "rejected", "invalid"])
def test_confirm_discovery_only_allows_pending_status(status: str) -> None:
    with _test_session() as db:
        db.add(Tenant(id="tenant_demo", name="Demo"))
        db.add(KnowledgeBase(id="kb_demo", tenant_id="tenant_demo", name="默认知识库"))
        suggestion = KnowledgeDiscoverySuggestion(
            tenant_id="tenant_demo",
            knowledge_base_id="kb_demo",
            document_id="doc_1",
            suggestion_type="warning",
            title="状态检查",
            status=status,
        )
        db.add(suggestion)
        db.commit()

        with pytest.raises(KnowledgeDiscoveryConflictError, match="只有待处理建议可以确认"):
            KnowledgeService(db).confirm_discovery(suggestion)

        db.refresh(suggestion)
        assert suggestion.status == status


def test_confirm_discovery_rolls_back_resource_when_status_commit_fails(monkeypatch: pytest.MonkeyPatch) -> None:
    with _test_session() as db:
        db.add(Tenant(id="tenant_demo", name="Demo"))
        db.add(KnowledgeBase(id="kb_demo", tenant_id="tenant_demo", name="默认知识库"))
        suggestion = KnowledgeDiscoverySuggestion(
            tenant_id="tenant_demo",
            knowledge_base_id="kb_demo",
            document_id="doc_1",
            suggestion_type="tool",
            title="会员权益核对",
            payload_json={"name": "member.benefit_reconcile", "url": "/api/mock/member"},
        )
        db.add(suggestion)
        db.commit()
        original_commit = db.commit
        monkeypatch.setattr(db, "commit", lambda: (_ for _ in ()).throw(RuntimeError("commit failed")))

        with pytest.raises(RuntimeError, match="commit failed"):
            KnowledgeService(db).confirm_discovery(suggestion)

        monkeypatch.setattr(db, "commit", original_commit)
        assert db.exec(select(Tool).where(Tool.name == "member.benefit_reconcile")).first() is None
        persisted = db.get(KnowledgeDiscoverySuggestion, suggestion.id)
        assert persisted is not None
        assert persisted.status == "pending"


def test_discovered_skill_rejects_unknown_fields_instead_of_dropping_them() -> None:
    payload = {
        "skill_id": "expense.travel_approval",
        "name": "差旅报销审批",
        "nodes": [
            {
                "node_id": "start",
                "type": "response",
                "name": "反馈结果",
                "instruction": "反馈结果。",
                "prompt": "这个字段不能被静默丢弃",
            }
        ],
        "edges": [],
        "start_node_id": "start",
        "terminal_node_ids": ["start"],
    }

    with pytest.raises(KnowledgeDiscoveryValidationError, match="nodes.0.*prompt"):
        validate_discovered_skill(payload)


def test_discovered_skill_requires_connected_start_to_terminal_graph() -> None:
    payload = {
        "skill_id": "expense.travel_approval",
        "name": "差旅报销审批",
        "nodes": [
            {
                "node_id": "start",
                "type": "collect_info",
                "name": "收集材料",
                "instruction": "收集材料。",
            },
            {
                "node_id": "reply",
                "type": "response",
                "name": "反馈结果",
                "instruction": "反馈结果。",
                "allowed_actions": ["answer_user"],
            },
            {
                "node_id": "orphan",
                "type": "response",
                "name": "孤立节点",
                "instruction": "处理孤立步骤。",
            },
        ],
        "edges": [{"source_node_id": "start", "next_node_id": "reply"}],
        "start_node_id": "start",
        "terminal_node_ids": ["reply"],
    }

    with pytest.raises(KnowledgeDiscoveryValidationError, match="无法从开始节点到达"):
        validate_discovered_skill(payload)

    payload["edges"].append({"source_node_id": "start", "next_node_id": "orphan"})
    with pytest.raises(KnowledgeDiscoveryValidationError, match="无法到达结束节点"):
        validate_discovered_skill(payload)


def test_discovery_only_marks_valid_skill_as_pending(monkeypatch: pytest.MonkeyPatch) -> None:
    valid_skill = {
        "skill_id": "expense.travel_approval",
        "name": "差旅报销审批",
        "nodes": [
            {
                "node_id": "collect",
                "type": "collect_info",
                "name": "收集材料",
                "instruction": "收集报销单据。",
            },
            {
                "node_id": "reply",
                "type": "response",
                "name": "反馈结果",
                "instruction": "反馈审批结果。",
                "allowed_actions": ["answer_user", "handoff_human"],
            },
        ],
        "edges": [{"source_node_id": "collect", "next_node_id": "reply"}],
        "start_node_id": "collect",
        "terminal_node_ids": ["reply"],
    }
    model_output = {
        "discoveries": [
            {
                "suggestion_type": "skill",
                "title": "差旅报销审批",
                "reason": "原文包含完整流程。",
                "payload": {"draft_skill": valid_skill},
            },
            {
                "suggestion_type": "skill",
                "title": "错误格式的技能",
                "reason": "模型使用了非标准字段。",
                "payload": {
                    "draft_skill": {
                        **valid_skill,
                        "nodes": [{"id": "collect", "label": "收集材料"}],
                    }
                },
            },
        ]
    }

    monkeypatch.setattr(LLMClient, "__init__", lambda self, model_config: None)
    monkeypatch.setattr(LLMClient, "generate_json", lambda self, prompt, payload: model_output)

    with _test_session() as db:
        db.add(Tenant(id="tenant_demo", name="Demo"))
        db.add(KnowledgeBase(id="kb_demo", tenant_id="tenant_demo", name="默认知识库"))
        db.add(
            ModelConfig(
                id="model_demo",
                tenant_id="tenant_demo",
                name="测试模型",
                api_key_encrypted="unused",
                model="test",
                is_default=True,
            )
        )
        document = KnowledgeDocument(
            tenant_id="tenant_demo",
            knowledge_base_id="kb_demo",
            filename="expense.md",
            file_type="md",
            title="报销制度",
        )
        bucket = KnowledgeBucket(
            tenant_id="tenant_demo",
            knowledge_base_id="kb_demo",
            document_id=document.id,
            bucket_key="expense",
            title="差旅报销",
            summary="差旅报销审批流程",
            metadata_json={"content": "提交材料后审批并反馈结果。"},
        )
        job = KnowledgeIngestJob(
            tenant_id="tenant_demo",
            knowledge_base_id="kb_demo",
            document_id=document.id,
            filename="expense.md",
            status="running",
            stage="discovering",
        )
        db.add(document)
        db.add(bucket)
        db.add(job)
        db.commit()

        KnowledgeService(db)._discover_from_document(  # noqa: SLF001
            "tenant_demo", "kb_demo", document, [bucket], job
        )

        rows = db.exec(select(KnowledgeDiscoverySuggestion)).all()
        assert {row.title: row.status for row in rows} == {
            "差旅报销审批": "pending",
            "错误格式的技能": "invalid",
        }
        valid_row = next(row for row in rows if row.status == "pending")
        stored_skill = valid_row.payload_json["draft_skill"]
        assert stored_skill["nodes"][0]["node_id"] == "collect"
        assert stored_skill["nodes"][0]["expected_user_info"] == []


def _b64(text: str) -> str:
    return base64.b64encode(text.encode("utf-8")).decode("ascii")


def _test_session():
    engine = create_engine(
        "sqlite://",
        connect_args={"check_same_thread": False},
        poolclass=StaticPool,
    )
    SQLModel.metadata.create_all(engine)
    return Session(engine)
